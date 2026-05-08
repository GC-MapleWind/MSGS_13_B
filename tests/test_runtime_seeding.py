import os
import subprocess
import sys
import tempfile
import time
import unittest
from pathlib import Path
from uuid import uuid4

import openpyxl
import psycopg


PROJECT_ROOT = Path(__file__).resolve().parents[1]
POSTGRES_IMAGE = "postgres:17-alpine"
POSTGRES_USER = "maplewind"
POSTGRES_PASSWORD = "maplewind"


def _run_python_snippet(code: str, env: dict[str, str]) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, "-c", code],
        cwd=PROJECT_ROOT,
        env=env,
        capture_output=True,
        text=True,
        check=False,
    )


def _run_seed_command(env: dict[str, str]) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, "-m", "scripts.seed_real_data"],
        cwd=PROJECT_ROOT,
        env=env,
        capture_output=True,
        text=True,
        check=False,
    )


def _run_docker_command(args: list[str]) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        ["docker", *args],
        capture_output=True,
        text=True,
        check=False,
    )


def _count_rows(database_url: str, table: str) -> int:
    with psycopg.connect(database_url) as conn:
        table_exists = conn.execute(
            """
            SELECT COUNT(*)
            FROM information_schema.tables
            WHERE table_schema = 'public' AND table_name = %s
            """,
            (table,),
        ).fetchone()[0]
        if table_exists == 0:
            return 0
        return conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]


def _create_seed_files(init_dir: Path) -> None:
    init_dir.mkdir(parents=True, exist_ok=True)

    roster_path = init_dir / "25-2 단풍바람 명부.xlsx"
    settlement_path = init_dir / "메생결산시트.xlsx"

    roster_wb = openpyxl.Workbook()
    roster_ws = roster_wb.active
    if roster_ws is None:
        raise RuntimeError("Failed to create roster worksheet")
    roster_ws.title = "25-2"
    roster_ws.append(
        [
            "name",
            "gender",
            "dept",
            "student_id",
            "nickname",
            "level",
            "server",
            "character_gender",
            "job",
        ]
    )
    roster_ws.append([
        "seed-user",
        "남자",
        "dev",
        20260001,
        "seednick",
        260,
        "LUNA",
        "",
        "Hero",
    ])
    roster_wb.save(roster_path)

    settlement_wb = openpyxl.Workbook()
    settlement_ws = settlement_wb.active
    if settlement_ws is None:
        raise RuntimeError("Failed to create settlement worksheet")
    settlement_ws.title = "Sheet1"
    settlement_ws.append(["id", "name", "nickname", "date", "img_name", "caption"])
    settlement_ws.append([1, "seed-user", "seednick", 260101, None, "Seed settlement"])
    settlement_wb.save(settlement_path)


class RuntimeSeedingTests(unittest.TestCase):
    postgres_container: str | None = None
    postgres_port: str | None = None
    admin_url: str | None = None

    @classmethod
    def setUpClass(cls) -> None:
        docker_probe = _run_docker_command(["version", "--format", "{{.Server.Version}}"])
        if docker_probe.returncode != 0:
            raise unittest.SkipTest(f"Docker is unavailable: {docker_probe.stderr.strip()}")

        container_name = f"maplewind-test-postgres-{uuid4().hex[:12]}"
        started = _run_docker_command(
            [
                "run",
                "-d",
                "--rm",
                "--name",
                container_name,
                "-e",
                f"POSTGRES_USER={POSTGRES_USER}",
                "-e",
                f"POSTGRES_PASSWORD={POSTGRES_PASSWORD}",
                "-e",
                "POSTGRES_DB=postgres",
                "-p",
                "127.0.0.1::5432",
                POSTGRES_IMAGE,
            ]
        )
        if started.returncode != 0:
            raise unittest.SkipTest(f"Could not start PostgreSQL test container: {started.stderr.strip()}")

        cls.postgres_container = container_name
        try:
            port_result = _run_docker_command(["port", container_name, "5432/tcp"])
            if port_result.returncode != 0:
                raise RuntimeError(port_result.stderr.strip())
            cls.postgres_port = port_result.stdout.rsplit(":", 1)[-1].strip()
            cls.admin_url = (
                f"postgresql://{POSTGRES_USER}:{POSTGRES_PASSWORD}"
                f"@127.0.0.1:{cls.postgres_port}/postgres"
            )
            deadline = time.monotonic() + 30
            while True:
                try:
                    with psycopg.connect(cls.admin_url, connect_timeout=1) as conn:
                        conn.execute("SELECT 1")
                    break
                except psycopg.OperationalError:
                    if time.monotonic() > deadline:
                        raise
                    time.sleep(0.5)
        except Exception:
            _run_docker_command(["rm", "-f", container_name])
            raise

    @classmethod
    def tearDownClass(cls) -> None:
        if cls.postgres_container:
            _run_docker_command(["rm", "-f", cls.postgres_container])

    def _database_urls(self) -> tuple[str, str]:
        if self.admin_url is None or self.postgres_port is None:
            raise RuntimeError("PostgreSQL test container did not start")

        db_name = f"maplewind_test_{uuid4().hex}"
        with psycopg.connect(self.admin_url, autocommit=True) as conn:
            conn.execute(f"CREATE DATABASE {db_name}")

        sync_url = (
            f"postgresql://{POSTGRES_USER}:{POSTGRES_PASSWORD}"
            f"@127.0.0.1:{self.postgres_port}/{db_name}"
        )
        async_url = (
            f"postgresql+asyncpg://{POSTGRES_USER}:{POSTGRES_PASSWORD}"
            f"@127.0.0.1:{self.postgres_port}/{db_name}"
        )
        return sync_url, async_url

    def _base_env(self, database_url: str, init_dir: Path) -> dict[str, str]:
        env = os.environ.copy()
        env.update(
            {
                "DATABASE_URL": database_url,
                "ADMIN_SESSION_SECRET": "test-admin-secret",
                "JWT_SECRET_KEY": "test-jwt-secret",
                "INIT_DATA_DIR": str(init_dir),
                "PYTHONPATH": str(PROJECT_ROOT),
            }
        )
        return env

    def test_startup_does_not_seed_without_explicit_command(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            sync_url, async_url = self._database_urls()
            init_dir = Path(temp_dir) / "init-data"
            env = self._base_env(async_url, init_dir)

            result = _run_python_snippet(
                (
                    "import asyncio\n"
                    "from src.database import init_db\n"
                    "from src.main import app\n"
                    "async def _run():\n"
                    "    await init_db()\n"
                    "    async with app.router.lifespan_context(app):\n"
                    "        pass\n"
                    "asyncio.run(_run())\n"
                ),
                env,
            )

            self.assertEqual(result.returncode, 0, msg=result.stderr)
            self.assertEqual(_count_rows(sync_url, "users"), 0)
            self.assertEqual(_count_rows(sync_url, "characters"), 0)
            self.assertEqual(_count_rows(sync_url, "comments"), 0)

    def test_src_main_has_no_scripts_import(self) -> None:
        src_main = (PROJECT_ROOT / "src" / "main.py").read_text(encoding="utf-8")
        self.assertNotIn("from scripts.", src_main)

    def test_explicit_seed_command_uses_init_data_dir(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            sync_url, async_url = self._database_urls()
            init_dir = Path(temp_dir) / "seed-input"
            _create_seed_files(init_dir)
            env = self._base_env(async_url, init_dir)

            result = _run_seed_command(env)

            self.assertEqual(result.returncode, 0, msg=result.stderr)
            self.assertEqual(_count_rows(sync_url, "users"), 1)
            self.assertEqual(_count_rows(sync_url, "characters"), 1)
            self.assertEqual(_count_rows(sync_url, "settlements"), 1)

            with psycopg.connect(sync_url) as conn:
                row = conn.execute(
                    "SELECT username, name FROM users WHERE name = %s", ("seed-user",)
                ).fetchone()
            self.assertIsNotNone(row)
            self.assertEqual(row[0], "20260001")

    def test_explicit_seed_command_fails_with_missing_files(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            _sync_url, async_url = self._database_urls()
            init_dir = Path(temp_dir) / "missing-seed-input"
            init_dir.mkdir(parents=True, exist_ok=True)
            env = self._base_env(async_url, init_dir)

            result = _run_seed_command(env)

            self.assertNotEqual(result.returncode, 0)


if __name__ == "__main__":
    unittest.main()
